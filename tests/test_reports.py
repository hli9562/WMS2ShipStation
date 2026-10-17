import io
import json
from datetime import datetime

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import text

from tests.test_imports import upload, workbook
from tests.test_purchases import FailingProvider, setup_order
from wms2ship.config import Settings
from wms2ship.main import create_app


def test_review_filters_snapshots_and_report_reconciliation(tmp_path):
    with TestClient(create_app(Settings(data_dir=tmp_path))) as client:
        assert client.get("/api/review").json()["total"] == 0
        order_id, headers, payload = setup_order(client)
        result = client.post(
            f"/api/orders/{order_id}/label-attempts", json=payload, headers=headers
        )
        assert result.status_code == 200
        failed_id, headers, payload = setup_order(client, "FAILED-ORDER")
        client.app.state.providers["fake"] = FailingProvider()
        client.post(f"/api/orders/{failed_id}/label-attempts", json=payload, headers=headers)
        upload(client, workbook(order="WAITING-ORDER"))
        with client.app.state.engine.begin() as db:
            db.execute(
                text("UPDATE orders SET origin_json=:origin WHERE id=:id"),
                {"origin": json.dumps({"name": "Changed after purchase"}), "id": order_id},
            )
            db.execute(
                text(
                    "UPDATE labels SET charge_minor=1234,currency='USD',tracking_number='000123456'"
                )
            )
            db.execute(text("UPDATE label_attempts SET created_at='2026-10-01T23:59:59Z'"))
        data = client.get("/api/review").json()
        assert data["summary"]["orders"] == 3
        assert data["summary"]["purchased"] == 1
        assert data["summary"]["exceptions"] == 1
        assert data["summary"]["not_started"] == 1
        assert data["summary"]["charges"] == [
            {"mode": "simulation", "currency": "USD", "minor": 1234}
        ]
        purchase = next(r for r in data["rows"] if r["state"] == "purchased")
        assert purchase["shipper"] == "Test Warehouse"
        assert client.get("/api/review?mode=live").json()["total"] == 0
        assert client.get("/api/review?state=failed").json()["total"] == 1
        assert client.get("/api/review?state=uncertain").json()["total"] == 0
        assert client.get("/api/review?search=000123456").json()["total"] == 1
        assert (
            client.get("/api/review?start=2026-10-01&end=2026-10-01&mode=simulation").json()[
                "total"
            ]
            == 2
        )
        assert client.get("/api/review?start=2026-10-02&mode=simulation").json()["total"] == 0
        assert client.get("/api/review?start=2026-10-02&end=2026-10-01").status_code == 422
        assert client.get("/api/review/report.xlsx?mode=invalid").status_code == 422
        response = client.get("/api/review/report.xlsx?state=purchased")
        assert response.status_code == 200
        book = load_workbook(io.BytesIO(response.content))
        assert book.sheetnames == ["Summary", "Shipments"]
        sheet = book["Shipments"]
        assert sheet.max_row == 2
        assert sheet["H2"].value == "Test Warehouse"
        assert sheet["L2"].value == "000123456"
        assert sheet["M2"].value == 12.34
        assert isinstance(sheet["C2"].value, datetime)
        assert sheet.freeze_panes == "A2"
        summary = dict(book["Summary"].iter_rows(values_only=True))
        assert summary["Purchased"] == 1
        assert summary["simulation / USD"] == 12.34
        with client.app.state.engine.begin() as db:
            db.execute(text("UPDATE labels SET charge_minor=NULL,currency=NULL"))
        assert client.get("/api/review").json()["summary"]["unknown_charges"] == 1


def test_report_includes_all_pages_and_treats_imported_text_as_literal(tmp_path):
    with TestClient(create_app(Settings(data_dir=tmp_path))) as client:
        with client.app.state.engine.begin() as db:
            for i in range(105):
                db.execute(
                    text("""INSERT INTO orders
                    (id,wms_order_no,platform_reference,created_at,updated_at)
                    VALUES (:id,:number,:ref,'2026-10-01T00:00:00Z','2026-10-01T00:00:00Z')"""),
                    {"id": str(i), "number": f"ORDER-{i:03}", "ref": "=1+1"},
                )
        data = client.get("/api/review").json()
        assert len(data["rows"]) == 100
        assert data["total"] == 105
        assert len(client.get("/api/review?offset=100").json()["rows"]) == 5
        book = load_workbook(io.BytesIO(client.get("/api/review/report.xlsx").content))
        sheet = book["Shipments"]
        assert sheet.max_row == 106
        assert sheet["B2"].value == "=1+1"
        assert sheet["B2"].data_type == "s"
        assert dict(book["Summary"].iter_rows(values_only=True))["Orders"] == 105
        empty = load_workbook(
            io.BytesIO(client.get("/api/review/report.xlsx?state=purchased").content)
        )
        assert empty["Shipments"].max_row == 1
