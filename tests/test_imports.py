import io
import json

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import text

from wms2ship.config import Settings
from wms2ship.main import create_app


def workbook(
    *,
    order="WMS-001",
    weight="0.460 kg",
    size="20 x 18.5 x 14 cm",
    quantity=2,
    count=1,
    tracking="",
    extra_pick=False,
    formula=False,
    duplicate=False,
):
    book = Workbook()
    sheet = book.active
    sheet.title = "出库单"
    sheet.append(
        [
            "Outbound Order No\n出库单号",
            "Recipient/收件人",
            "Country/Region/国家/地区",
            "Address1/地址1",
            "City/城市名称",
            "Post code/邮编",
            "Package Qty/包裹数量",
            "Package 1 size",
            "Package 1 Package Weight",
            "SKU 1 SKU",
            "Province Code/省州编码",
            "Package 1 Tracking No.",
        ]
    )
    row = [
        order,
        "Synthetic Recipient",
        "Canada",
        "100 Test Street",
        "Test City",
        "T0A 0A0",
        count,
        size,
        weight,
        "TEST-SKU",
        "AB",
        tracking,
    ]
    sheet.append(row)
    if duplicate:
        sheet.append(row)
    if formula:
        sheet.cell(2, 9, "=1+1")
    pick = book.create_sheet("拣货库位明细")
    pick.append(["出库单号", "SKU", "出库数量"])
    pick.append([order, "TEST-SKU", quantity])
    if extra_pick:
        pick.append([order, "TEST-SKU", 1])
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def upload(client, payload, filename="test.xlsx"):
    return client.post(
        f"/api/imports?filename={filename}",
        content=payload,
        headers={"x-csrf-token": client.app.state.csrf_token},
    )


def test_sample_import_duplicate_and_provenance(tmp_path):
    with TestClient(create_app(Settings(data_dir=tmp_path))) as client:
        result = upload(client, workbook()).json()
        assert result["created"] == 1
        rows = client.get(f"/api/imports/{result['id']}").json()["rows"]
        assert {row["sheet"] for row in rows} == {"出库单", "拣货库位明细"}
        order_id = client.get("/api/orders").json()[0]["id"]
        order = client.get(f"/api/orders/{order_id}").json()
        assert order["items"][0]["quantity"] == 2
        package = order["packages"][0]
        assert [package[k] for k in ("length_mm", "width_mm", "height_mm", "weight_g")] == [
            200,
            185,
            140,
            460,
        ]
        assert package["weight_oz"] == 17
        assert package["length_in"] == 8
        assert package["width_in"] == 8
        assert package["height_in"] == 6
        assert order["recipient"]["country"] == "CA"
        # Changed re-import must preserve operator's configuration and source data.
        with client.app.state.engine.begin() as db:
            db.execute(
                text("UPDATE orders SET shipping_json=:value WHERE id=:id"),
                {"value": json.dumps({"carrier": "configured"}), "id": order_id},
            )
        result2 = upload(client, workbook(weight="1 kg"), "second.xlsx").json()
        assert result2["duplicates"] == 1
        order2 = client.get(f"/api/orders/{order_id}").json()
        assert order2["shipping"] == {"carrier": "configured"}
        assert order2["packages"][0]["weight_g"] == 460
        with client.app.state.engine.connect() as db:
            assert db.execute(text("SELECT count(*) FROM label_attempts")).scalar() == 0


def test_invalid_rows_and_workbook_rejections(tmp_path):
    with TestClient(create_app(Settings(data_dir=tmp_path))) as client:
        for payload in (
            workbook(count=2),
            workbook(weight="0.460"),
            workbook(extra_pick=True),
            workbook(duplicate=True),
            workbook(quantity=0),
        ):
            response = upload(client, payload)
            assert response.status_code == 201
            assert response.json()["invalid"] > 0
            assert response.json()["created"] == 0
        assert upload(client, workbook(formula=True)).status_code == 422
        assert upload(client, workbook(order=123456)).status_code == 422
        assert upload(client, b"not excel").status_code == 422
        assert upload(client, workbook(), "old.xls").status_code == 400
        assert client.get("/api/orders").json() == []


def test_existing_tracking_warning_and_upload_limit(tmp_path):
    settings = Settings(data_dir=tmp_path)
    with TestClient(create_app(settings)) as client:
        result = upload(client, workbook(tracking="EXISTING-TEST")).json()
        assert result["warnings"] == 1
        detail = client.get(f"/api/imports/{result['id']}").json()
        assert any("Existing WMS tracking" in " ".join(row["issues"]) for row in detail["rows"])
        settings.max_upload_bytes = 20
        assert upload(client, workbook()).status_code == 413


def test_bad_headers_and_row_limit(tmp_path):
    settings = Settings(data_dir=tmp_path, max_sheet_rows=1)
    with TestClient(create_app(settings)) as client:
        assert upload(client, workbook(duplicate=True)).status_code == 422
        book = Workbook()
        sheet = book.active
        sheet.title = "出库单"
        sheet.append(["出库单号", "出库单号"])
        stream = io.BytesIO()
        book.save(stream)
        response = upload(client, stream.getvalue())
        assert response.status_code == 422
        assert "Ambiguous" in response.json()["detail"]


def test_concurrent_reimports_create_one_order(tmp_path):
    from concurrent.futures import ThreadPoolExecutor

    with TestClient(create_app(Settings(data_dir=tmp_path))) as client:
        payload = workbook()
        with ThreadPoolExecutor(max_workers=2) as pool:
            responses = list(pool.map(lambda _: upload(client, payload), range(2)))
        assert all(response.status_code == 201 for response in responses)
        assert sum(response.json()["created"] for response in responses) == 1
        assert sum(response.json()["duplicates"] for response in responses) == 1
        assert len(client.get("/api/orders").json()) == 1


def test_original_export_header_and_size_conventions(tmp_path):
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(workbook()))
    sheet = book["出库单"]
    sheet.cell(1, 8, r"Package 1\nPackage Size/包裹尺寸")
    sheet.cell(1, 9, r"Package 1\nPackage Weight/包裹重量")
    sheet.cell(1, 10, r"SKU 1\nSKU")
    sheet.cell(1, 12, r"Package 1\nTracking No./物流跟踪号")
    sheet.cell(2, 8, "20.000*18.500*14.000*cm")
    sheet.cell(2, 9, "0.460kg")
    pick = book["拣货库位明细"]
    pick.cell(1, 2, r"SKU 1\nSKU")
    pick.cell(1, 3, r"SKU 1\nOutbound Qty/出库数量")
    stream = io.BytesIO()
    book.save(stream)
    with TestClient(create_app(Settings(data_dir=tmp_path))) as client:
        result = upload(client, stream.getvalue()).json()
        assert result["created"] == 1
        assert result["invalid"] == 0


@pytest.mark.parametrize("weight,expected_oz", [("1 lb", 16), ("16 oz", 16), ("453.59237 g", 16)])
def test_import_us_units_round_up_each_dimension(tmp_path, weight, expected_oz):
    with TestClient(create_app(Settings(data_dir=tmp_path))) as client:
        response = upload(client, workbook(weight=weight, size="10.25 x 8.5 x 4 in"))
        assert response.json()["created"] == 1, response.json()
        order_id = client.get("/api/orders").json()[0]["id"]
        package = client.get(f"/api/orders/{order_id}").json()["packages"][0]
        assert package["weight_oz"] == expected_oz
        assert [package[k] for k in ("length_in", "width_in", "height_in")] == [11, 9, 4]
        assert package["weight_g"] == 454
        # Re-import preserves the already converted package.
        assert upload(client, workbook(weight="2 kg")).json()["duplicates"] == 1
        again = client.get(f"/api/orders/{order_id}").json()["packages"][0]
        assert again["weight_oz"] == expected_oz


def test_older_imports_automatically_expose_us_units(tmp_path):
    with TestClient(create_app(Settings(data_dir=tmp_path))) as client:
        assert upload(client, workbook()).json()["created"] == 1
        order_id = client.get("/api/orders").json()[0]["id"]
        with client.app.state.engine.begin() as db:
            db.execute(text(
                "UPDATE packages SET weight_oz=NULL,length_in=NULL,width_in=NULL,height_in=NULL"
            ))
        package = client.get(f"/api/orders/{order_id}").json()["packages"][0]
        assert package["weight_oz"] == 17
        assert package["length_in"] == 8
        assert package["width_in"] == 8
        assert package["height_in"] == 6
