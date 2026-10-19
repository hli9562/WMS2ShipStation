import io
from contextlib import contextmanager
from datetime import date

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from wms2ship.carrier_fees import shipping_name
from wms2ship.config import Settings
from wms2ship.main import create_app
from wms2ship.shipment_source import ShipmentSource
from wms2ship.shipstation import ProviderUnavailable


def shipment(identity=1, **changes):
    return {
        "shipmentId": identity,
        "orderId": identity,
        "orderNumber": f"ORDER-{identity}",
        "shipDate": "2026-10-01",
        "createDate": "2026-09-30T12:00:00",
        "shipmentCost": 10.01,
        "insuranceCost": 0,
        "trackingNumber": "000123456",
        "carrierCode": "usps",
        "serviceCode": "usps_ground_advantage",
        "voided": False,
        "shipTo": {"name": "Synthetic Recipient", "country": "US"},
        **changes,
    }


@contextmanager
def account_client(tmp_path, handler):
    settings = Settings(
        data_dir=tmp_path, SHIPSTATION_API_KEY="mock-key", SHIPSTATION_API_SECRET="mock-secret"
    )
    with TestClient(create_app(settings)) as client:
        client.app.state.shipment_source.close()
        client.app.state.shipment_source = ShipmentSource(
            "mock-key", "mock-secret", transport=httpx.MockTransport(handler)
        )
        yield client


def test_account_fees_without_local_orders_and_carrier_column_excel(tmp_path):
    shipments = [
        shipment(1),
        shipment(2, shipmentCost=20.02),
        shipment(3, shipmentCost=None),
        shipment(4, orderNumber="=ZERO", shipmentCost=0),
        shipment(5, voided=True),
        shipment(6, testLabel=True),
        shipment(7, shipDate="2026-10-02"),
    ]
    calls = []

    def handler(request):
        calls.append(request)
        assert request.method == "GET" and request.url.path == "/shipments"
        assert request.url.params["shipDateStart"] == "2026-10-01"
        assert request.url.params["shipDateEnd"] == "2026-10-01"
        return httpx.Response(200, json={"shipments": shipments, "total": len(shipments)})

    with account_client(tmp_path, handler) as client:
        assert client.get("/api/orders").json() == []
        query = "?start=2026-10-01&end=2026-10-01"
        data = client.get("/api/carrier-fees" + query).json()
        assert data["total"] == data["summary"]["orders"] == 4
        assert data["summary"]["source"] == "ShipStation account shipments"
        assert data["summary"]["date_basis"] == "Ship date"
        assert data["summary"]["missing_fees"] == 1
        assert data["summary"]["voided_excluded"] == 1
        assert data["summary"]["totals"] == [{"currency": "USD", "charge_minor": 3003}]
        response = client.get("/api/carrier-fees/report.xlsx" + query)
        book = load_workbook(io.BytesIO(response.content))
        assert book.sheetnames == ["Shipping Fees"]
        sheet = book.active
        assert sheet.max_column == 5 and sheet.max_row == 6
        assert list(next(sheet.values)) == [
            "Order #",
            "Tracking #",
            "Recipient",
            "Carrier",
            "Shipping Fee (USD)",
        ]
        by_order = {row[0].value: row for row in list(sheet.rows)[1:-1]}
        assert by_order["ORDER-3"][4].value is None
        assert by_order["=ZERO"][0].data_type == "s"
        assert by_order["=ZERO"][4].value == 0
        assert by_order["ORDER-1"][1].value == "000123456"
        assert by_order["ORDER-1"][3].value == "USPS\nUSPS Ground Advantage"
        assert by_order["ORDER-1"][3].alignment.wrap_text
        assert sheet.row_dimensions[2].height == 32
        assert sheet.auto_filter.ref == "A1:E5"
        assert by_order["ORDER-1"][4].number_format == '"$"#,##0.00'
        assert sheet["E6"].value == 30.03 and "incomplete" in sheet["A6"].value
        assert len(calls) == 1  # Page and Excel reuse the same complete account snapshot.
        assert client.get("/api/orders").json() == []
        assert client.get("/api/carrier-fees?start=2026-10-02&end=2026-10-01").status_code == 422


def test_insurance_pagination_and_export_total(tmp_path):
    records = [
        shipment(i, orderId=1, shipmentCost="0.10", insuranceCost="0.02") for i in range(1, 106)
    ]
    calls = []

    def handler(request):
        calls.append(request)
        page = int(request.url.params["page"])
        return httpx.Response(
            200, json={"shipments": records[(page - 1) * 100 : page * 100], "total": 105}
        )

    with account_client(tmp_path, handler) as client:
        query = "?start=2026-10-01&end=2026-10-01"
        data = client.get("/api/carrier-fees" + query).json()
        assert data["summary"]["orders"] == 1 and data["total"] == 105
        assert len(data["rows"]) == 100
        assert data["summary"]["totals"] == [{"currency": "USD", "charge_minor": 1260}]
        assert len(client.get("/api/carrier-fees" + query + "&offset=100").json()["rows"]) == 5
        book = load_workbook(
            io.BytesIO(client.get("/api/carrier-fees/report.xlsx" + query).content)
        )
        assert book.active.max_row == 107
        assert book.active["A107"].value == "TOTAL" and book.active["E107"].value == 12.6
        assert len(calls) == 2


@pytest.mark.parametrize("status", [401, 403, 429, 500])
def test_remote_failures_never_return_zero_or_partial_fee_totals(tmp_path, status):
    with account_client(tmp_path, lambda request: httpx.Response(status)) as client:
        for endpoint in ("/api/carrier-fees", "/api/carrier-fees/report.xlsx"):
            response = client.get(endpoint)
            assert response.status_code == 502
            assert "detail" in response.json() and "summary" not in response.json()
            assert "mock-secret" not in response.text


def test_incomplete_pages_and_offset_limit_do_not_produce_a_bill():
    for response in (
        {"shipments": [], "total": 1},
        {"shipments": [shipment()], "total": 10001},
        {"shipments": [{"orderNumber": "missing id"}], "total": 1},
    ):
        source = ShipmentSource(
            "mock",
            "mock",
            transport=httpx.MockTransport(lambda request: httpx.Response(200, json=response)),
        )
        try:
            with pytest.raises(ProviderUnavailable):
                source.shipments(date(2026, 10, 1), date(2026, 10, 1))
        finally:
            source.close()


def test_duplicate_shipments_are_counted_once(tmp_path):
    one = shipment(1)
    with account_client(
        tmp_path,
        lambda request: httpx.Response(
            200, json={"shipments": [one, one, shipment(2)], "total": 2}
        ),
    ) as client:
        data = client.get("/api/carrier-fees?start=2026-10-01&end=2026-10-01").json()
        assert data["total"] == 2
        assert data["summary"]["totals"] == [{"currency": "USD", "charge_minor": 2002}]


def test_missing_configuration_is_explicit_and_empty_account_has_zero_total(tmp_path):
    with TestClient(create_app(Settings(data_dir=tmp_path))) as client:
        assert client.get("/api/carrier-fees").status_code == 502
    with account_client(
        tmp_path, lambda request: httpx.Response(200, json={"shipments": [], "total": 0})
    ) as client:
        data = client.get("/api/carrier-fees").json()
        assert data["total"] == 0 and data["summary"]["missing_fees"] == 0


@pytest.mark.parametrize(
    "code,name",
    [
        ("usps_ground_advantage", "USPS Ground Advantage"),
        ("ups_2nd_day_air", "UPS 2nd Day Air"),
        ("fedex_2day", "FedEx 2Day"),
    ],
)
def test_specific_service_names(code, name):
    assert shipping_name(code) == name


def test_api_supplied_service_name_is_preserved(tmp_path):
    data = shipment(serviceName="UPS 2nd Day Air A.M.", carrierName="UPS")
    with account_client(
        tmp_path, lambda request: httpx.Response(200, json={"shipments": [data], "total": 1})
    ) as client:
        query = "?start=2026-10-01&end=2026-10-01"
        row = client.get("/api/carrier-fees" + query).json()["rows"][0]
        assert row["service_name"] == "UPS 2nd Day Air A.M."
        book = load_workbook(
            io.BytesIO(client.get("/api/carrier-fees/report.xlsx" + query).content)
        )
        assert book.active["D2"].value == "UPS\nUPS 2nd Day Air A.M."
