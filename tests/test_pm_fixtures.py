import io
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import text

from wms2ship.config import Settings
from wms2ship.db import make_engine, migrate
from wms2ship.importer import import_workbook

FIXTURES = Path(__file__).parent / "fixtures" / "pm"


def test_committed_pm_workbooks_have_expected_acceptance_results(tmp_path):
    paths = sorted(FIXTURES.glob("*.xlsx"))
    assert len(paths) == 5
    assert paths[0].read_bytes() == paths[3].read_bytes()
    settings = Settings(data_dir=tmp_path)
    migrate(settings)
    engine = make_engine(settings)
    expected = [(1, 0, 0), (1, 0, 0), (3, 0, 0), (0, 1, 0), (0, 0, 1)]
    try:
        for path, counts in zip(paths, expected, strict=True):
            payload = path.read_bytes()
            book = load_workbook(io.BytesIO(payload))
            assert book.sheetnames == ["出库单", "拣货库位明细"]
            assert not any(cell.data_type == "f" for sheet in book for row in sheet for cell in row)
            result = import_workbook(engine, payload, path.name, settings)
            assert tuple(result[key] for key in ("created", "duplicates", "invalid")) == counts
        with engine.connect() as db:
            assert db.execute(text("SELECT count(*) FROM orders")).scalar() == 5
            assert db.execute(text("SELECT count(*) FROM label_attempts")).scalar() == 0
            assert (
                db.execute(
                    text(
                        "SELECT json_extract(recipient_json,'$.postal') FROM orders "
                        "WHERE wms_order_no='PM-US-001'"
                    )
                ).scalar()
                == "02108"
            )
            assert (
                db.execute(
                    text(
                        "SELECT json_extract(recipient_json,'$.country') FROM orders "
                        "WHERE wms_order_no='PM-CA-001'"
                    )
                ).scalar()
                == "CA"
            )
    finally:
        engine.dispose()
