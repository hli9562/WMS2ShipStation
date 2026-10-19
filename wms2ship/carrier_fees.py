"""Account-wide ShipStation shipment charges for customer billing."""

import io
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from wms2ship.ship_dates import SHIPPER_TIMEZONE
from wms2ship.shipstation import ProviderUnavailable

router = APIRouter(prefix="/api/carrier-fees")


def fee_filters(start: date | None = None, end: date | None = None):
    today = datetime.now(SHIPPER_TIMEZONE).date()
    start = start or today.replace(day=1)
    end = end or today
    if start > end:
        raise HTTPException(422, "From date must be on or before Through date.")
    if end == date.max:
        raise HTTPException(422, "Through date must be earlier than 9999-12-31.")
    return {"start": start, "end": end}


def amount_minor(value):
    try:
        amount = Decimal(str(value))
        cents = amount * 100
        if not amount.is_finite() or amount < 0 or cents != cents.to_integral_value():
            return None
        return int(cents)
    except (InvalidOperation, ValueError, TypeError):
        return None


def fee_data(request, selected):
    source = request.app.state.shipment_source
    if source is None:
        raise ProviderUnavailable(
            "ShipStation V1 API credentials are not configured on the server. "
            "Configure API Key and API Secret to read account shipments. No total was calculated."
        )
    shipments = source.shipments(selected["start"], selected["end"])
    rows = []
    totals = defaultdict(int)
    groups = {}
    voided = 0
    for shipment in shipments:
        if shipment.get("testLabel") is True or shipment.get("isTestLabel") is True:
            continue
        try:
            day = date.fromisoformat(str(shipment["shipDate"])[:10])
        except (KeyError, ValueError):
            raise ProviderUnavailable(
                "A ShipStation shipment has an invalid ship date. No total was calculated."
            ) from None
        if not selected["start"] <= day <= selected["end"]:
            continue
        if shipment.get("voided") is True:
            voided += 1
            continue
        postage = amount_minor(shipment.get("shipmentCost"))
        insurance = amount_minor(shipment.get("insuranceCost", 0))
        charge = None if postage is None or insurance is None else postage + insurance
        recipient = shipment.get("shipTo") or {}
        shipment_id = str(shipment["shipmentId"])
        number = shipment.get("orderNumber") or ""
        row = dict(
            label_id=shipment_id,
            order_id=str(shipment.get("orderId") or number or f"shipment:{shipment_id}"),
            wms_order_no=number or f"Shipment {shipment_id}",
            platform_reference="",
            purchased_at=shipment.get("createDate") or "",
            ship_date=day.isoformat(),
            carrier=shipment.get("carrierCode") or "Not recorded",
            service=shipment.get("serviceCode") or "Not recorded",
            tracking_number=shipment.get("trackingNumber") or "",
            recipient=recipient.get("name") or "",
            country=recipient.get("country") or "",
            charge_minor=charge,
            postage_minor=postage,
            insurance_minor=insurance,
            currency="USD",
            void_status="active",
            note="Return label" if shipment.get("isReturnLabel") is True else "",
        )
        rows.append(row)
        key = (row["carrier"], row["currency"] or "Unknown")
        group = groups.setdefault(
            key, dict(carrier=key[0], currency=key[1], labels=0, charge_minor=0, missing_fees=0)
        )
        group["labels"] += 1
        if row["charge_minor"] is None:
            group["missing_fees"] += 1
        else:
            totals[row["currency"]] += row["charge_minor"]
            group["charge_minor"] += row["charge_minor"]
    if not rows:
        totals["USD"] = 0
    return rows, dict(
        orders=len({r["order_id"] for r in rows}),
        labels=len(rows),
        missing_fees=sum(r["charge_minor"] is None for r in rows),
        voided_excluded=voided,
        source="ShipStation account shipments",
        date_basis="Ship date",
        totals=[dict(currency=c, charge_minor=n) for c, n in sorted(totals.items())],
        carriers=[groups[key] for key in sorted(groups)],
    )


@router.get("")
def fee_summary(
    request: Request,
    selected: dict = Depends(fee_filters),
    offset: int = Query(0, ge=0),
    limit: int = Query(100, ge=1, le=200),
):
    rows, summary = fee_data(request, selected)
    return dict(
        filters=selected,
        timezone=str(SHIPPER_TIMEZONE),
        summary=summary,
        rows=rows[offset : offset + limit],
        total=len(rows),
    )


def fee_workbook(rows, summary, selected):
    book = Workbook()
    sheet = book.active
    sheet.title = "Shipping Fees"
    sheet.append(["Order #", "Tracking #", "Recipient", "Shipping Fee (USD)"])
    for row in rows:
        sheet.append(
            [
                row["wms_order_no"],
                row["tracking_number"],
                row["recipient"],
                None if row["charge_minor"] is None else Decimal(row["charge_minor"]) / 100,
            ]
        )
        sheet.cell(sheet.max_row, 4).number_format = '"$"#,##0.00'
    total = sum(r["charge_minor"] for r in rows if r["charge_minor"] is not None)
    title = (
        "TOTAL"
        if not summary["missing_fees"]
        else (f"TOTAL (incomplete: {summary['missing_fees']} missing fees)")
    )
    sheet.append([title, None, None, Decimal(total) / 100])
    sheet.cell(sheet.max_row, 4).number_format = '"$"#,##0.00'
    for row in (sheet[1], sheet[sheet.max_row]):
        for cell in row:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="082E58")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:D{max(1, sheet.max_row - 1)}"
    for column, width in zip("ABCD", (32, 32, 35, 25), strict=True):
        sheet.column_dimensions[column].width = width
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_LETTER
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
    sheet.print_options.horizontalCentered = True
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


@router.get("/report.xlsx")
def download_fees(request: Request, selected: dict = Depends(fee_filters)):
    rows, summary = fee_data(request, selected)
    filename = f"carrier-fees-{selected['start']}-{selected['end']}.xlsx"
    return Response(
        fee_workbook(rows, summary, selected),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
