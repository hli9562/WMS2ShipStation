"""Manager review and XLSX export of stored label activity (not delivery tracking)."""

import io
import json
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal
from typing import Literal

from fastapi import APIRouter, Depends, HTTPException, Query, Request
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import text

from wms2ship.providers import provider_mode

router = APIRouter(prefix="/api/review")


def filters(
    start: date | None = None,
    end: date | None = None,
    state: Literal[
        "all", "not_started", "prepared", "sending", "purchased", "failed", "uncertain", "removed"
    ] = "all",
    mode: Literal["all", "live", "simulation", "test", "unassigned"] = "all",
    search: str = Query("", max_length=100),
):
    if start and end and start > end:
        raise HTTPException(422, "Start date must be on or before end date.")
    return dict(start=start, end=end, state=state, mode=mode, search=search.strip())


def review_data(request, selected):
    with request.app.state.engine.connect() as db:
        records = (
            db.execute(
                text("""
            SELECT o.id AS order_id, o.wms_order_no, o.platform_reference,
              COALESCE(a.created_at,o.created_at) AS created_at,
              CASE WHEN l.void_status='voided' THEN 'removed'
                ELSE COALESCE(a.state,'not_started') END AS state,
              a.state AS purchase_state, a.id AS attempt_id,
              a.error_code, a.request_snapshot_json, o.recipient_json, o.origin_json,
              o.shipping_json, p.provider_code, l.id AS label_id, l.tracking_number,
              l.charge_minor, l.currency, l.void_status, l.void_updated_at
            FROM orders o LEFT JOIN label_attempts a ON a.order_id=o.id
            LEFT JOIN provider_accounts p ON p.id=a.provider_account_id
            LEFT JOIN labels l ON l.attempt_id=a.id
            ORDER BY COALESCE(a.created_at,o.created_at) DESC,o.id,a.id,l.id
        """)
            )
            .mappings()
            .all()
        )
    rows = []
    for record in records:
        r = dict(record)
        snapshot = json.loads(r.pop("request_snapshot_json") or "{}")
        recipient = snapshot.get("recipient") or json.loads(r.pop("recipient_json") or "{}")
        origin = snapshot.get("origin") or json.loads(r.pop("origin_json") or "{}")
        shipping = snapshot.get("shipping") or json.loads(r.pop("shipping_json") or "{}")
        # Remove current-order JSON even when the historical snapshot was used.
        for key in ("recipient_json", "origin_json", "shipping_json"):
            r.pop(key, None)
        r.update(
            mode=provider_mode(r["provider_code"]) if r["attempt_id"] else "unassigned",
            recipient=recipient.get("name", ""),
            country=recipient.get("country", ""),
            shipper=origin.get("name", ""),
            carrier=shipping.get("carrier_code", ""),
            service=shipping.get("service_code", ""),
            ship_date=shipping.get("ship_date", ""),
        )
        activity_date = date.fromisoformat(r["created_at"][:10])
        if selected["start"] and activity_date < selected["start"]:
            continue
        if selected["end"] and activity_date > selected["end"]:
            continue
        if selected["state"] != "all" and r["state"] != selected["state"]:
            continue
        if selected["mode"] != "all" and r["mode"] != selected["mode"]:
            continue
        if (
            selected["search"]
            and selected["search"].casefold()
            not in " ".join(
                str(r.get(k) or "")
                for k in (
                    "wms_order_no",
                    "platform_reference",
                    "tracking_number",
                    "recipient",
                    "shipper",
                )
            ).casefold()
        ):
            continue
        rows.append(r)
    counts = Counter(r["state"] for r in rows)
    charges = defaultdict(int)
    for r in rows:
        if r["purchase_state"] == "purchased" and r["charge_minor"] is not None:
            charges[(r["mode"], r["currency"])] += r["charge_minor"]
    summary = dict(
        records=len(rows),
        orders=len({r["order_id"] for r in rows}),
        purchased=sum(r["purchase_state"] == "purchased" for r in rows),
        not_started=counts["not_started"],
        exceptions=sum(
            r["state"] in ("failed", "uncertain") or bool(r["error_code"]) for r in rows
        ),
        states=dict(counts),
        charges=[dict(mode=m, currency=c, minor=n) for (m, c), n in sorted(charges.items())],
        unknown_charges=sum(
            r["purchase_state"] == "purchased" and r["charge_minor"] is None for r in rows
        ),
    )
    return rows, summary


@router.get("")
def dashboard(
    request: Request,
    selected: dict = Depends(filters),
    offset: int = Query(0, ge=0),
    limit: int = Query(100, ge=1, le=200),
):
    rows, summary = review_data(request, selected)
    return dict(summary=summary, rows=rows[offset : offset + limit], total=len(rows))


COLUMNS = [
    ("wms_order_no", "WMS order"),
    ("platform_reference", "Platform reference"),
    ("created_at", "Activity time (UTC)"),
    ("state", "Label status"),
    ("mode", "Mode"),
    ("recipient", "Recipient"),
    ("country", "Destination"),
    ("shipper", "Shipper"),
    ("carrier", "Carrier"),
    ("service", "Service"),
    ("ship_date", "Ship date"),
    ("tracking_number", "Tracking number"),
    ("charge_minor", "Charge"),
    ("currency", "Currency"),
    ("void_status", "Void status"),
    ("void_updated_at", "Void updated (UTC)"),
    ("error_code", "Exception"),
    ("attempt_id", "Attempt ID"),
    ("label_id", "Label ID"),
]


def report_workbook(rows, summary, selected):
    book = Workbook()
    overview = book.active
    overview.title = "Summary"
    overview.append(["Shipment review report", "Value"])
    overview.append(["Generated (UTC)", datetime.now(timezone.utc).replace(tzinfo=None)])
    overview["B2"].number_format = "yyyy-mm-dd hh:mm:ss"
    overview.append(
        ["Source", "Local WMS2ShipStation database; stored label activity, not delivery status"]
    )
    for key, value in selected.items():
        overview.append([f"Filter: {key}", str(value) if value is not None else "All dates"])
    for key in ("records", "orders", "purchased", "not_started", "exceptions", "unknown_charges"):
        overview.append([key.replace("_", " ").title(), summary[key]])
    overview.append(["Charges by mode and currency", "Amount"])
    for charge in summary["charges"]:
        overview.append(
            [f"{charge['mode']} / {charge['currency']}", Decimal(charge["minor"]) / 100]
        )
        overview.cell(overview.max_row, 2).number_format = "#,##0.00"
    overview.append(
        [
            "Reading the detail",
            "One row per label attempt, or unstarted order. Retries appear separately.",
        ]
    )
    overview.append(
        [
            "Simulation",
            "Simulation and test labels are not valid postage; test charges are not live spend.",
        ]
    )
    sheet = book.create_sheet("Shipments")
    sheet.append([title for _, title in COLUMNS])
    for row in rows:
        values = []
        for key, _ in COLUMNS:
            value = row.get(key)
            if key == "charge_minor" and value is not None:
                value = Decimal(value) / 100
            if key == "created_at":
                value = (
                    datetime.fromisoformat(value.replace("Z", "+00:00"))
                    .astimezone(timezone.utc)
                    .replace(tzinfo=None)
                )
            if key == "ship_date" and value:
                value = date.fromisoformat(value)
            values.append(value)
        sheet.append(values)
        sheet.cell(sheet.max_row, 3).number_format = "yyyy-mm-dd hh:mm:ss"
        sheet.cell(sheet.max_row, 11).number_format = "yyyy-mm-dd"
        sheet.cell(sheet.max_row, 13).number_format = "#,##0.00"
    sheet.auto_filter.ref = sheet.dimensions
    for ws in book:
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="082E58")
        for column in ws.columns:
            letter = get_column_letter(column[0].column)
            ws.column_dimensions[letter].width = min(
                48, max(18, max(len(str(c.value or "")) for c in column) + 2)
            )
            for cell in column:
                if isinstance(cell.value, str):
                    # Imported text is always literal, including leading '=' and URLs.
                    cell.data_type = "s"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "1:1"
    overview.column_dimensions["A"].width = 34
    overview.column_dimensions["B"].width = 90
    for row in overview:
        overview.row_dimensions[row[0].row].height = 32
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


@router.get("/report.xlsx")
def download_report(request: Request, selected: dict = Depends(filters)):
    rows, summary = review_data(request, selected)
    return Response(
        report_workbook(rows, summary, selected),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="shipment-review.xlsx"'},
    )
