"""Strict XLWMS import: workbook contents are data, never executable instructions."""

import hashlib
import io
import json
import re
import zipfile
from datetime import UTC, datetime
from decimal import ROUND_CEILING, Decimal, InvalidOperation
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy import MetaData, Table, select

from wms2ship.config import Settings

PARSER_VERSION = "2"


class ImportRejected(ValueError):
    pass


def norm(value):
    return re.sub(r"[^\w]", "", str(value or "").replace(r"\n", " ").lower()).replace("_", "")


# English, Chinese, or a bilingual header separated by whitespace/slashes.
FIELDS = {
    "order": ("Outbound Order No", "出库单号"),
    "platform": ("Platform Number", "平台单号"),
    "name": ("Recipient", "收件人"),
    "phone": ("Telephone", "电话"),
    "email": ("Email", "邮箱"),
    "company": ("Company", "公司名称"),
    "country": ("Country/Region", "国家/地区"),
    "state": ("Province Code", "省州编码"),
    "state_name": ("Province Name", "省州名称"),
    "city": ("City", "城市名称"),
    "postal": ("Post code", "邮编"),
    "house": ("House No.", "门牌号"),
    "address1": ("Address1", "地址1"),
    "address2": ("Address2", "地址2"),
    "district": ("District", "区县"),
    "package_qty": ("Package Qty", "包裹数量"),
    "size": ("Package 1 size", "包裹1尺寸", "Package 1 Package Size/包裹尺寸"),
    "length": ("Package 1 length", "包裹1长", "Package 1 Package Size/包裹长"),
    "width": ("Package 1 width", "包裹1宽", "Package 1 Package Size/包裹宽"),
    "height": ("Package 1 height", "包裹1高", "Package 1 Package Size/包裹高"),
    "weight": ("Package 1 Package Weight", "包裹1重量", "Package 1 Package Weight/包裹重量"),
    "tracking": ("Package 1 Tracking No.", "物流跟踪号"),
    "sku": ("SKU 1 SKU", "SKU1"),
    "description": ("SKU 1 Product Name", "商品名称", "SKU 1 Product Name/产品名称"),
}
PICK_FIELDS = {
    "order": ("Outbound Order No", "出库单号"),
    "sku": ("SKU", "SKU", "SKU 1 SKU"),
    "quantity": ("Outbound Qty", "出库数量", "SKU 1 Outbound Qty/出库数量"),
}


def header_map(cells, fields, required):
    result = {}
    headers = [norm(c.value) for c in cells]
    for field, aliases in fields.items():
        accepted = {norm(a) for a in aliases} | {
            norm("".join(aliases[:2])),
            norm("".join(reversed(aliases[:2]))),
        }
        matches = [i for i, value in enumerate(headers) if value in accepted]
        if len(matches) > 1:
            raise ImportRejected(f"Ambiguous header for {aliases[0]}.")
        if matches:
            result[field] = matches[0]
        elif field in required:
            raise ImportRejected(f"Missing required header: {aliases[0]} / {aliases[1]}.")
    return result


def read_sheet(workbook, name, fields, required, settings):
    if name not in workbook.sheetnames:
        raise ImportRejected(f"Missing worksheet: {name}.")
    sheet = workbook[name]
    if sheet.max_column and sheet.max_column > 128:
        raise ImportRejected(f"{name}: too many columns (maximum 128).")
    rows = sheet.iter_rows(max_row=settings.max_sheet_rows + 2, max_col=128)
    headers = next(rows, ())
    mapping = header_map(headers, fields, required)
    records = []
    for number, cells in enumerate(rows, 2):
        if not any(cell.value is not None for cell in cells):
            continue
        if number > settings.max_sheet_rows + 1:
            raise ImportRejected(f"{name}: exceeds {settings.max_sheet_rows} data rows.")
        if any(cell.data_type == "f" for cell in cells):
            raise ImportRejected(f"{name}, row {number}: formulas are not supported.")
        values = {key: cells[index].value for key, index in mapping.items()}
        # Reject numeric IDs: leading zeros/precision may already have been lost.
        for key in ("order", "sku", "postal", "phone", "tracking", "platform"):
            value = values.get(key)
            if value is not None and not isinstance(value, str):
                raise ImportRejected(f"{name}, row {number}: {key} must be stored as Excel text.")
        records.append((number, values))
    return records


def text_value(values, key, required=False):
    value = str(values.get(key) or "").strip()
    if required and not value:
        raise ValueError(f"{key} is required.")
    return value


def decimal_value(value):
    try:
        number = Decimal(str(value))
    except InvalidOperation:
        raise ValueError("Invalid numeric measurement.") from None
    if not number.is_finite() or number <= 0:
        raise ValueError("Measurements and quantities must be positive.")
    return number


def exact_integer(number):
    if number != number.to_integral_value():
        raise ValueError("Unsupported measurement or quantity precision.")
    return int(number)


def measurements(values):
    size = text_value(values, "size", required=True).lower().replace("×", "x").replace("*", "x")
    match = re.fullmatch(r"\s*([\d.]+)\s*x\s*([\d.]+)\s*x\s*([\d.]+)\s*x?\s*(cm|mm|in)\s*", size)
    if not match:
        raise ValueError("Package size must include length x width x height and cm, mm, or in.")
    factor = {"cm": Decimal(10), "mm": Decimal(1), "in": Decimal("25.4")}[match[4]]
    result = {}
    for key, raw in zip(("length", "width", "height"), match.groups()[:3], strict=True):
        value = decimal_value(raw)
        separate = values.get(key)
        if separate is not None and decimal_value(separate) != value:
            raise ValueError(f"Combined size conflicts with separate {key}.")
        millimeters = value * factor
        result[f"{key}_mm"] = int(millimeters.to_integral_value(rounding=ROUND_CEILING))
        result[f"{key}_in"] = int(
            (millimeters / Decimal("25.4")).to_integral_value(rounding=ROUND_CEILING)
        )
    weight = re.fullmatch(
        r"\s*([\d.]+)\s*(kg|g|lb|oz)\s*", text_value(values, "weight", required=True).lower()
    )
    if not weight:
        raise ValueError("Package weight must include kg, g, lb, or oz.")
    factor = {
        "kg": Decimal(1000),
        "g": Decimal(1),
        "lb": Decimal("453.59237"),
        "oz": Decimal("28.349523125"),
    }[weight[2]]
    grams = decimal_value(weight[1]) * factor
    result["weight_g"] = int(grams.to_integral_value(rounding=ROUND_CEILING))
    result["weight_oz"] = int(
        (grams / Decimal("28.349523125")).to_integral_value(rounding=ROUND_CEILING)
    )
    return result


def parse_workbook(payload, settings):
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            if sum(info.file_size for info in archive.infolist()) > settings.max_uncompressed_bytes:
                raise ImportRejected("Workbook exceeds the uncompressed size limit.")
            if len(archive.infolist()) > 1000:
                raise ImportRejected("Workbook contains too many archive entries.")
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=False)
    except (zipfile.BadZipFile, OSError, KeyError, ValueError) as exc:
        if isinstance(exc, ImportRejected):
            raise
        raise ImportRejected("Cannot read this file as an unencrypted XLSX workbook.") from None
    try:
        outbound = read_sheet(
            workbook,
            "出库单",
            FIELDS,
            {
                "order",
                "name",
                "country",
                "address1",
                "city",
                "postal",
                "package_qty",
                "size",
                "weight",
                "sku",
            },
            settings,
        )
        picking = read_sheet(
            workbook, "拣货库位明细", PICK_FIELDS, {"order", "sku", "quantity"}, settings
        )
        if not outbound:
            raise ImportRejected("The outbound worksheet has no orders.")
        return outbound, picking
    finally:
        workbook.close()


def prepare_orders(outbound, picking):
    quantities = {}
    for _, row in picking:
        key = (text_value(row, "order"), text_value(row, "sku"))
        quantities.setdefault(key, []).append(row.get("quantity"))
    prepared = []
    seen = set()
    for row_number, source in outbound:
        issues = []
        order_no = text_value(source, "order")
        shipment = None
        try:
            order_no = text_value(source, "order", True)
            if order_no in seen:
                raise ValueError(
                    "Repeated outbound order rows require review; package mapping is ambiguous."
                )
            seen.add(order_no)
            if decimal_value(source.get("package_qty")) != 1:
                raise ValueError("Only one unambiguous package per order is currently supported.")
            sku = text_value(source, "sku", True)
            matches = quantities.get((order_no, sku), [])
            if len(matches) != 1:
                raise ValueError("SKU quantity join is missing or ambiguous; review picking rows.")
            if any(key[0] == order_no and key[1] != sku for key in quantities):
                raise ValueError("Multiple SKU mapping requires review before import.")
            country = text_value(source, "country", True).upper()
            country = {
                "CANADA": "CA",
                "加拿大": "CA",
                "UNITED STATES": "US",
                "USA": "US",
                "美国": "US",
            }.get(country, country)
            if country not in {"US", "CA"}:
                raise ValueError("Only US and Canada destinations are supported.")
            if text_value(source, "house") or text_value(source, "district"):
                raise ValueError(
                    "Populated house/district fields require address-composition review."
                )
            if text_value(source, "state_name") and not text_value(source, "state"):
                raise ValueError(
                    "A province/state code is required when a province name is supplied."
                )
            recipient = {
                k: text_value(source, k, k in {"name", "address1", "city", "postal"})
                for k in (
                    "name",
                    "phone",
                    "email",
                    "company",
                    "address1",
                    "address2",
                    "city",
                    "state",
                    "postal",
                )
            }
            recipient["country"] = country
            package = measurements(source)
            package["source_tracking"] = text_value(source, "tracking") or None
            if package["source_tracking"]:
                issues.append(
                    "Existing WMS tracking number: review before purchasing another label."
                )
            shipment = {
                "order_no": order_no,
                "platform": text_value(source, "platform"),
                "recipient": recipient,
                "package": package,
                "sku": sku,
                "description": text_value(source, "description"),
                "quantity": exact_integer(decimal_value(matches[0])),
            }
        except ValueError as exc:
            issues.append(str(exc))
        prepared.append(
            {
                "row": row_number,
                "source": source,
                "order_no": order_no,
                "shipment": shipment,
                "issues": issues,
            }
        )
    # All rows of a repeated source order are unsafe, including the first one.
    counts = {}
    for item in prepared:
        counts[item["order_no"]] = counts.get(item["order_no"], 0) + 1
    for item in prepared:
        if counts[item["order_no"]] > 1:
            item["shipment"] = None
            item["issues"] = ["Repeated outbound order rows require package-mapping review."]
    return prepared


def import_workbook(engine, payload, filename, settings: Settings):
    outbound, picking = parse_workbook(payload, settings)
    prepared = prepare_orders(outbound, picking)
    metadata = MetaData()
    tables = {
        name: Table(name, metadata, autoload_with=engine)
        for name in ("import_batches", "import_rows", "orders", "packages", "order_items")
    }
    now = datetime.now(UTC).isoformat()
    batch_id = str(uuid4())
    summary = {"created": 0, "duplicates": 0, "invalid": 0, "warnings": 0}
    linked = {}
    with engine.connect() as db:
        # SQLite local single-writer reservation also serializes concurrent imports.
        db.exec_driver_sql("BEGIN IMMEDIATE")
        try:
            db.execute(
                tables["import_batches"]
                .insert()
                .values(
                    id=batch_id,
                    original_filename=filename,
                    sha256=hashlib.sha256(payload).hexdigest(),
                    parser_version=PARSER_VERSION,
                    created_at=now,
                    summary_json="{}",
                )
            )
            for item in prepared:
                shipment = item["shipment"]
                order_id = None
                if shipment is None:
                    summary["invalid"] += 1
                else:
                    order_id = db.execute(
                        select(tables["orders"].c.id).where(
                            tables["orders"].c.wms_order_no == shipment["order_no"]
                        )
                    ).scalar_one_or_none()
                    if order_id:
                        summary["duplicates"] += 1
                        item["issues"].append(
                            "Existing order retained; compare this source row "
                            "before changing stored data."
                        )
                    else:
                        order_id = str(uuid4())
                        db.execute(
                            tables["orders"]
                            .insert()
                            .values(
                                id=order_id,
                                wms_order_no=shipment["order_no"],
                                platform_reference=shipment["platform"],
                                recipient_json=json.dumps(
                                    shipment["recipient"], ensure_ascii=False
                                ),
                                created_at=now,
                                updated_at=now,
                            )
                        )
                        db.execute(
                            tables["packages"]
                            .insert()
                            .values(
                                id=str(uuid4()),
                                order_id=order_id,
                                source_package_key="1",
                                **shipment["package"],
                            )
                        )
                        db.execute(
                            tables["order_items"]
                            .insert()
                            .values(
                                id=str(uuid4()),
                                order_id=order_id,
                                source_item_key="1",
                                sku=shipment["sku"],
                                description=shipment["description"],
                                quantity=shipment["quantity"],
                            )
                        )
                        summary["created"] += 1
                    linked[item["order_no"]] = order_id
                if item["issues"]:
                    summary["warnings"] += 1
                db.execute(
                    tables["import_rows"]
                    .insert()
                    .values(
                        id=str(uuid4()),
                        batch_id=batch_id,
                        sheet_name="出库单",
                        row_number=item["row"],
                        source_order_no=item["order_no"],
                        order_id=order_id,
                        evidence_json=json.dumps(item["source"], ensure_ascii=False, default=str),
                        issues_json=json.dumps(item["issues"], ensure_ascii=False),
                    )
                )
            for number, row in picking:
                db.execute(
                    tables["import_rows"]
                    .insert()
                    .values(
                        id=str(uuid4()),
                        batch_id=batch_id,
                        sheet_name="拣货库位明细",
                        row_number=number,
                        source_order_no=text_value(row, "order"),
                        order_id=linked.get(text_value(row, "order")),
                        evidence_json=json.dumps(row, ensure_ascii=False, default=str),
                        issues_json="[]",
                    )
                )
            db.execute(
                tables["import_batches"]
                .update()
                .where(tables["import_batches"].c.id == batch_id)
                .values(summary_json=json.dumps(summary))
            )
            db.commit()
        except Exception:
            db.rollback()
            raise
    return {"id": batch_id, **summary}
